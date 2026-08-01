from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "Bao-cao-do-an-UDA-Assessment-Hub.docx"
LOGO = ROOT / "assets" / "uda-logo.png"
GREEN = "018F41"; DARK = "01823C"; ORANGE = "F58220"; INK = "263238"; LIGHT = "F4F6F9"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
sec.header_distance = sec.footer_distance = Inches(.492)

def font(run, size=11, bold=False, color=INK, italic=False):
    run.font.name = "Calibri"; run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri"); run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size); run.bold = bold; run.italic = italic; run.font.color.rgb = RGBColor.from_string(color)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"; normal.font.size = Pt(11); normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_before = Pt(0); normal.paragraph_format.space_after = Pt(8); normal.paragraph_format.line_spacing = 1.333
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
for name, size, before, after, color in [("Heading 1",16,18,10,GREEN),("Heading 2",13,12,6,DARK),("Heading 3",12,8,4,DARK)]:
    s=styles[name]; s.font.name="Calibri"; s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor.from_string(color)
    s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after); s.paragraph_format.keep_with_next=True

def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement("w:shd"); shd.set(qn("w:fill"),fill); tcPr.append(shd)

def set_cell_width(cell, dxa):
    tcPr=cell._tc.get_or_add_tcPr(); tcW=tcPr.find(qn("w:tcW"))
    if tcW is None:
        tcW=OxmlElement("w:tcW"); tcPr.append(tcW)
    tcW.set(qn("w:w"),str(dxa)); tcW.set(qn("w:type"),"dxa")

def table(headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False; t.style="Table Grid"
    widths=widths or [9360//len(headers)]*len(headers)
    tblPr=t._tbl.tblPr; tblW=tblPr.find(qn("w:tblW")); tblW.set(qn("w:w"),"9360"); tblW.set(qn("w:type"),"dxa")
    ind=OxmlElement("w:tblInd"); ind.set(qn("w:w"),"120"); ind.set(qn("w:type"),"dxa"); tblPr.append(ind)
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=str(h); shade(c,GREEN); set_cell_width(c,widths[i]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in c.paragraphs[0].runs: font(r,10,bold=True,color="FFFFFF")
    for row in rows:
        cells=t.add_row().cells
        for i,val in enumerate(row):
            cells[i].text=str(val); set_cell_width(cells[i],widths[i]); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after=Pt(3); p.paragraph_format.line_spacing=1.15
                for r in p.runs: font(r,9.5)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

def bullets(items):
    for item in items:
        p=doc.add_paragraph(style="List Bullet"); p.paragraph_format.left_indent=Inches(.375); p.paragraph_format.first_line_indent=Inches(-.194); p.paragraph_format.space_after=Pt(4); p.paragraph_format.line_spacing=1.208
        p.add_run(item)

def numbered(items):
    numbering = doc.part.numbering_part.element
    existing_abs = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    existing_num = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abs_id = max(existing_abs, default=0) + 1; num_id = max(existing_num, default=0) + 1
    abstract = OxmlElement("w:abstractNum"); abstract.set(qn("w:abstractNumId"), str(abs_id))
    multi = OxmlElement("w:multiLevelType"); multi.set(qn("w:val"), "singleLevel"); abstract.append(multi)
    lvl = OxmlElement("w:lvl"); lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start"); start.set(qn("w:val"), "1"); lvl.append(start)
    fmt = OxmlElement("w:numFmt"); fmt.set(qn("w:val"), "decimal"); lvl.append(fmt)
    txt = OxmlElement("w:lvlText"); txt.set(qn("w:val"), "%1."); lvl.append(txt)
    suff = OxmlElement("w:suff"); suff.set(qn("w:val"), "space"); lvl.append(suff)
    ppr = OxmlElement("w:pPr"); ind = OxmlElement("w:ind"); ind.set(qn("w:left"), "540"); ind.set(qn("w:hanging"), "280"); ppr.append(ind); lvl.append(ppr)
    abstract.append(lvl); numbering.append(abstract)
    num = OxmlElement("w:num"); num.set(qn("w:numId"), str(num_id)); aid = OxmlElement("w:abstractNumId"); aid.set(qn("w:val"), str(abs_id)); num.append(aid); numbering.append(num)
    for item in items:
        p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.375); p.paragraph_format.first_line_indent=Inches(-.194); p.paragraph_format.space_after=Pt(4); p.paragraph_format.line_spacing=1.208; p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        numPr=OxmlElement("w:numPr"); ilvl=OxmlElement("w:ilvl"); ilvl.set(qn("w:val"),"0"); nid=OxmlElement("w:numId"); nid.set(qn("w:val"),str(num_id)); numPr.append(ilvl); numPr.append(nid); p._p.get_or_add_pPr().append(numPr)
        p.add_run(item)

def p(text, bold_lead=None):
    par=doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        font(par.add_run(bold_lead),bold=True,color=DARK); par.add_run(text[len(bold_lead):])
    else: par.add_run(text)
    return par

def page_break(): doc.add_page_break()

# Running furniture
h=sec.header.paragraphs[0]; h.alignment=WD_ALIGN_PARAGRAPH.LEFT; font(h.add_run("ĐẠI HỌC ĐÔNG Á  |  UDA ASSESSMENT HUB"),9,bold=True,color=GREEN)
f=sec.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.CENTER; font(f.add_run("Báo cáo đồ án · Phiên bản 1.0 · 08/2026"),8,color="6B7280")

# Editorial cover (narrative_proposal preset + UDA named color override)
for _ in range(3): doc.add_paragraph()
if LOGO.exists():
    lp=doc.add_paragraph(); lp.alignment=WD_ALIGN_PARAGRAPH.CENTER; lp.add_run().add_picture(str(LOGO),width=Inches(3.7))
kp=doc.add_paragraph(); kp.alignment=WD_ALIGN_PARAGRAPH.CENTER; font(kp.add_run("BÁO CÁO ĐỒ ÁN CÔNG NGHỆ THÔNG TIN"),11,bold=True,color=ORANGE)
tp=doc.add_paragraph(); tp.alignment=WD_ALIGN_PARAGRAPH.CENTER; tp.paragraph_format.space_before=Pt(12); tp.paragraph_format.space_after=Pt(10); font(tp.add_run("UDA ASSESSMENT HUB"),28,bold=True,color=DARK)
sp=doc.add_paragraph(); sp.alignment=WD_ALIGN_PARAGRAPH.CENTER; font(sp.add_run("Nền tảng học tập, nhập đề bằng AI và quản trị khảo thí số"),15,bold=True,color=GREEN)
doc.add_paragraph()
table(["Hạng mục","Thông tin"],[
    ("Đơn vị định hướng","Trường Đại học Đông Á · Thành phố Đà Nẵng"),
    ("Loại sản phẩm","Ứng dụng web quản lý học, ôn tập, kiểm tra và thi chính thức"),
    ("Phiên bản báo cáo","1.0 · Tháng 08/2026"),
    ("Mã nguồn","github.com/Lo75d1/kahoot-web"),
],[2400,6960])
cp=doc.add_paragraph(); cp.alignment=WD_ALIGN_PARAGRAPH.CENTER; cp.paragraph_format.space_before=Pt(36); font(cp.add_run("ĐÀ NẴNG · 2026"),11,bold=True,color=DARK)

page_break()
doc.add_heading("TÓM TẮT",1)
p("UDA Assessment Hub là nền tảng web được phát triển từ nhu cầu số hóa toàn bộ vòng đời câu hỏi và kỳ thi tại cơ sở giáo dục đại học. Hệ thống kết hợp nhập đề bán tự động bằng AI, ngân hàng đề, học–ôn tập hiện đại, quiz tương tác realtime và một phân hệ thi chính thức có phân quyền, kiểm duyệt độc lập, niêm phong, kiểm tra điều kiện dự thi, chấm điểm phía máy chủ, công bố kết quả, phúc khảo và nhật ký truy vết.")
p("Điểm khác biệt của sản phẩm so với công cụ quiz giải trí là các quy tắc kiểm soát được đặt tại tầng cơ sở dữ liệu bằng Row Level Security và các hàm nghiệp vụ bảo mật. Vì vậy, việc đổi giao diện hoặc gọi API trực tiếp không thể bỏ qua các nguyên tắc maker–checker, quyền theo vai trò và trạng thái của đề thi. Phiên bản hiện tại được xây dựng với Next.js 16, React 19, Supabase/PostgreSQL và triển khai trên Vercel.")
doc.add_heading("Từ khóa",2); p("AI nhập đề; ngân hàng câu hỏi; khảo thí số; maker–checker; Supabase; Next.js; học thích ứng; kiểm tra trực tuyến.")

doc.add_heading("MỤC LỤC TÓM LƯỢC",1)
numbered(["Bối cảnh và bài toán","Mục tiêu, phạm vi và đối tượng","Khảo sát và yêu cầu hệ thống","Giải pháp và kiến trúc","Thiết kế chức năng","An toàn và quản trị khảo thí","Hiện thực và kiểm thử","Đánh giá, giới hạn và hướng phát triển","Kết luận và tài liệu tham khảo"])

page_break(); doc.add_heading("1. BỐI CẢNH VÀ BÀI TOÁN",1)
doc.add_heading("1.1 Bối cảnh",2)
p("Chuyển đổi số giáo dục không chỉ là đưa bài kiểm tra lên web. Một hệ thống dùng trong trường đại học phải trả lời được ai tạo đề, ai kiểm duyệt, đề nào được phép sử dụng, sinh viên nào đủ điều kiện, dữ liệu nào được gửi xuống trình duyệt, ai công bố điểm và mọi thay đổi sau chấm thi đã được phê duyệt ra sao.")
doc.add_heading("1.2 Vấn đề cần giải quyết",2)
bullets(["Tài liệu câu hỏi từ nhiều nguồn có cấu trúc không đồng nhất, khiến nhập liệu thủ công chậm và dễ sai.","Công cụ quiz phổ thông tối ưu cho tương tác nhưng thiếu quy trình kiểm duyệt nghiêm ngặt cho thi chính thức.","Đáp án hoặc logic chấm có thể bị lộ nếu toàn bộ dữ liệu được gửi xuống trình duyệt.","Điều kiện dự thi, biên bản giám thị, đổi điểm và phúc khảo thường nằm rời rạc ở nhiều tệp/hệ thống.","Nhà trường cần nhật ký truy vết và phân quyền theo chức trách để kiểm tra sau sự kiện."])
doc.add_heading("1.3 Định hướng UDA",2)
p("Giao diện sử dụng nhận diện xanh của Đại học Đông Á, logo chính thức và ngôn ngữ nghiệp vụ gần với giảng viên, sinh viên, phòng đào tạo, khảo thí và đảm bảo chất lượng. Thiết kế hướng đến khả năng trở thành một mô-đun trong hệ sinh thái số của trường, không tuyên bố là hệ thống chính thức khi chưa được nhà trường nghiệm thu.")

doc.add_heading("2. MỤC TIÊU, PHẠM VI VÀ ĐỐI TƯỢNG",1)
doc.add_heading("2.1 Mục tiêu",2)
numbered(["Rút ngắn thời gian nhập và chuẩn hóa đề từ văn bản/PDF/DOCX/CSV với sự hỗ trợ của AI.","Hỗ trợ học, ôn tập, kiểm tra nhanh và thi đấu trực tiếp trên cùng kho nội dung.","Thiết lập chu trình đề thi chính thức có phản biện, phê duyệt và niêm phong độc lập.","Kiểm soát ca thi, danh sách đủ điều kiện, bài làm, công bố điểm, phúc khảo và đổi điểm.","Cung cấp nền tảng có thể triển khai đám mây, kiểm thử tự động và mở rộng tích hợp LMS/SSO."])
doc.add_heading("2.2 Phạm vi phiên bản",2)
table(["Trong phạm vi","Ngoài phạm vi hiện tại"],[
    ("Web responsive; tài khoản; ngân hàng đề; AI nhập đề; QR; realtime; học/ôn; khảo thí", "Ứng dụng native iOS/Android"),
    ("RLS, RPC, audit, chấm server, phúc khảo và maker–checker đổi điểm", "Khóa trình duyệt cấp hệ điều hành và nhận diện khuôn mặt"),
    ("Triển khai Vercel + Supabase", "SSO sản xuất, tích hợp SIS/LMS chính thức và kiểm thử tải quy mô toàn trường"),
],[4680,4680])

doc.add_heading("3. KHẢO SÁT VÀ YÊU CẦU HỆ THỐNG",1)
doc.add_heading("3.1 Nhóm người dùng",2)
table(["Vai trò","Trách nhiệm chính"],[
    ("Sinh viên","Học, ôn tập, tham gia quiz, dự thi, xem kết quả và phúc khảo"),
    ("Giảng viên","Tạo/nhập đề, quản lý nội dung và đề nghị đổi điểm"),
    ("Trưởng bộ môn","Phản biện nội dung, chuẩn đầu ra và đáp án"),
    ("Khảo thí","Niêm phong, lập/mở/đóng ca, công bố kết quả"),
    ("Giám thị","Theo dõi lượt thi và lập biên bản"),
    ("Phòng đào tạo","Điều kiện dự thi và điều hành kỳ thi"),
    ("Đảm bảo chất lượng","Giám sát, xử lý phúc khảo và duyệt đổi điểm"),
    ("Quản trị viên","Quản trị kỹ thuật và phân quyền"),
],[2400,6960])
doc.add_heading("3.2 Yêu cầu phi chức năng",2)
bullets(["Bảo mật theo nguyên tắc quyền tối thiểu và không tin cậy dữ liệu phía khách.","Tính toàn vẹn của đề sau khi gửi duyệt và niêm phong.","Khả năng truy vết quyết định nghiệp vụ bằng audit log.","Responsive, thao tác rõ ràng và phản hồi trạng thái tự lưu.","Khả năng triển khai lặp lại bằng migration và kiểm tra build tự động."])

page_break(); doc.add_heading("4. GIẢI PHÁP VÀ KIẾN TRÚC",1)
doc.add_heading("4.1 Kiến trúc tổng thể",2)
p("Client Next.js/React cung cấp một ứng dụng đa chế độ. Route Handler xử lý trích xuất tài liệu và AI. Supabase đảm nhiệm Auth, PostgreSQL, RLS, RPC và Realtime. Vercel phân phối giao diện và API serverless. Cách tách này giảm vận hành máy chủ riêng nhưng vẫn giữ luật nghiệp vụ quan trọng ở cơ sở dữ liệu.")
table(["Lớp","Công nghệ","Chức năng"],[
    ("Presentation","Next.js 16, React 19, Tailwind CSS","Giao diện học, thi và quản trị"),
    ("Application","TypeScript, Route Handlers","Điều phối nhập tài liệu/AI và nghiệp vụ client"),
    ("Data/Security","Supabase PostgreSQL, RLS, RPC","Lưu trữ, phân quyền, chấm thi và audit"),
    ("Realtime","Supabase Realtime","Phòng chơi, người tham gia và kết quả trực tiếp"),
    ("Delivery","GitHub, Vercel","Quản lý phiên bản và triển khai"),
],[1800,2700,4860])
doc.add_heading("4.2 Mô hình dữ liệu trọng tâm",2)
table(["Nhóm","Bảng chính"],[
    ("Nội dung","quizzes, question_reviews, courses, departments"),
    ("Danh tính","profiles, auth.users"),
    ("Kỳ thi","exam_sessions, exam_eligibility, exam_attempts, exam_events"),
    ("Sau thi","grade_changes, appeals, audit_logs"),
    ("Học/realtime","classes, assignments, learning_progress, live_rooms, live_players"),
],[2400,6960])

doc.add_heading("5. THIẾT KẾ CHỨC NĂNG",1)
doc.add_heading("5.1 Nhập đề thông minh",2)
p("Người dùng có thể dán nội dung tự do, nhập theo marker, CSV hoặc quy tắc tùy chỉnh; tài liệu PDF/DOCX được trích xuất trước khi phân tích. AI được dùng như bộ hỗ trợ biến đổi cấu trúc, không thay thế khâu kiểm tra chuyên môn. Bộ đề luôn có bước xem lại/chỉnh sửa trước khi đưa vào quy trình chính thức.")
doc.add_heading("5.2 Học và ôn tập hiện đại",2)
bullets(["Luyện tập chủ động theo câu hỏi và phản hồi tức thời.","Ôn lỗi sai và lặp lại nội dung chưa vững.","Flashcard/tự kiểm tra để tăng khả năng nhớ lại chủ động.","Theo dõi tiến độ theo bộ đề và lớp học.","Quiz trực tiếp bằng mã PIN/QR để tăng tương tác trên lớp."])
doc.add_heading("5.3 Vòng đời đề thi",2)
table(["Trạng thái","Ý nghĩa","Chủ thể"],[
    ("draft","Đang biên soạn, được phép chỉnh sửa","Giảng viên"),
    ("in_review","Đã chốt snapshot/hash, chờ phản biện","Trưởng bộ môn/đơn vị duyệt"),
    ("changes_requested","Trả lại kèm nhận xét","Giảng viên sửa và gửi lại"),
    ("approved","Đạt yêu cầu chuyên môn","Người duyệt độc lập"),
    ("sealed","Đề được khóa để lập ca","Khảo thí độc lập"),
    ("retired","Ngừng sử dụng","Đơn vị có thẩm quyền"),
],[1800,4380,3180])
doc.add_heading("5.4 Ca thi và bài làm",2)
numbered(["Lập ca từ đề đã niêm phong, cấu hình thời gian và kiểm soát.","Nạp danh sách đủ điều kiện dự thi và lý do loại trừ.","Mở ca; sinh viên được xác thực trước khi tạo lượt thi duy nhất.","Máy chủ trả payload đã bỏ đáp án, giải thích và gợi ý.","Client tự lưu; các sự kiện rời tab/kết nối lại được ghi nhận.","Máy chủ chấm điểm, giữ kín kết quả tới khi cán bộ có quyền công bố."])

page_break(); doc.add_heading("6. AN TOÀN VÀ QUẢN TRỊ KHẢO THÍ",1)
doc.add_heading("6.1 Kiểm soát theo vai trò",2)
p("RLS giới hạn dòng dữ liệu người dùng được đọc/ghi; mọi thao tác nhạy cảm đi qua hàm RPC SECURITY DEFINER có kiểm tra vai trò, trạng thái và quan hệ giữa các chủ thể. Việc khóa luật ở tầng dữ liệu giúp giảm nguy cơ bỏ qua kiểm soát bằng cách gọi API trực tiếp.")
doc.add_heading("6.2 Maker–checker",2)
bullets(["Người sở hữu đề không thể tự phê duyệt.","Người niêm phong không được trùng người ra đề hoặc người duyệt.","Người đề nghị đổi điểm không thể tự duyệt đề nghị đó.","Kết luận phúc khảo yêu cầu vai trò khảo thí/đảm bảo chất lượng/quản trị."])
doc.add_heading("6.3 Bảo vệ đề và đáp án",2)
p("Trước khi gửi duyệt, hệ thống tạo SHA-256 từ tiêu đề, mô tả và nội dung câu hỏi. Khi phản biện, hash được tính lại để phát hiện thay đổi. Sau niêm phong, trigger cơ sở dữ liệu chặn sửa/xóa. Trong ca thi, RPC tạo payload an toàn bằng cách loại trường đúng/sai, giải thích, gợi ý và nguồn tham khảo; điểm được tính trong PostgreSQL.")
doc.add_heading("6.4 Nhật ký và xử lý sau thi",2)
table(["Sự kiện","Dữ liệu lưu","Mục đích"],[
    ("Bắt đầu/tự lưu/nộp","Lượt thi, tác nhân, thời gian","Khôi phục và xác minh tiến trình"),
    ("Chuyển tab/kết nối lại","Loại sự kiện, thời gian","Tín hiệu hỗ trợ giám sát"),
    ("Biên bản giám thị","Nội dung, mức độ, người lập","Hồ sơ sự cố"),
    ("Phúc khảo/đổi điểm","Lý do, quyết định, người duyệt","Trách nhiệm giải trình"),
    ("Audit nghiệp vụ","Before/after và metadata","Kiểm tra sau sự kiện"),
],[2100,3300,3960])

doc.add_heading("7. HIỆN THỰC VÀ KIỂM THỬ",1)
doc.add_heading("7.1 Phạm vi hiện thực",2)
table(["Hạng mục","Trạng thái"],[
    ("Nhập văn bản, PDF/DOCX, CSV/quy tắc, hỗ trợ AI","Hoàn thành"),
    ("Kho 65 bộ đề mẫu / 348 câu hỏi","Hoàn thành"),
    ("Học, ôn tập, lớp học và tiến độ","Hoàn thành mức đồ án"),
    ("Live quiz, mã PIN và QR","Hoàn thành"),
    ("8 vai trò và quy trình kiểm duyệt đề","Hoàn thành"),
    ("Thi chính thức, tự lưu, chấm server, công bố","Hoàn thành lõi"),
    ("Điều kiện dự thi, giám thị, phúc khảo, đổi điểm","Hoàn thành lõi"),
],[6500,2860])
doc.add_heading("7.2 Kết quả kiểm tra kỹ thuật",2)
table(["Kiểm tra","Kết quả tại thời điểm lập báo cáo"],[
    ("ESLint","Đạt, không có lỗi"),
    ("Vitest","22/22 kiểm thử đạt trên 9 tệp"),
    ("Next.js production build","Biên dịch thành công"),
    ("Migration Supabase 002–004","Áp dụng thành công"),
    ("GitHub/Vercel","Nhánh main được đẩy và tự động triển khai"),
],[4300,5060])
doc.add_heading("7.3 Kịch bản nghiệm thu chính",2)
bullets(["Không cho người ra đề tự duyệt đề của mình.","Không cho sửa nội dung sau khi gửi duyệt/niêm phong.","Sinh viên ngoài danh sách không thể tạo lượt thi.","Payload thi không chứa đáp án đúng.","Điểm chưa xuất hiện trước thao tác công bố.","Người đề nghị không thể tự duyệt đổi điểm.","Phúc khảo và biên bản tạo audit có dấu thời gian."])

doc.add_heading("8. ĐÁNH GIÁ, GIỚI HẠN VÀ HƯỚNG PHÁT TRIỂN",1)
doc.add_heading("8.1 Giá trị đạt được",2)
p("Sản phẩm hợp nhất trải nghiệm học tập linh hoạt với quy trình khảo thí có kiểm soát. Thiết kế phân tách rõ quiz luyện tập và bài thi chính thức; giảm nhập liệu nhờ AI nhưng vẫn đặt con người chịu trách nhiệm; đồng thời tạo nền móng dữ liệu đủ để kiểm tra và xử lý khiếu nại.")
doc.add_heading("8.2 Giới hạn",2)
bullets(["Chấm tự động hiện tối ưu cho câu hỏi lựa chọn; tự luận cần quy trình chấm tay/rubric.","Chưa tích hợp SSO và dữ liệu sinh viên chính thức của trường.","Tín hiệu chuyển tab chỉ hỗ trợ giám sát, không phải bằng chứng gian lận độc lập.","Chưa kiểm thử tải quy mô toàn trường và chưa có cam kết SLA.","Cần thẩm định pháp lý, bảo vệ dữ liệu cá nhân và quy chế nội bộ trước vận hành thật."])
doc.add_heading("8.3 Lộ trình",2)
table(["Giai đoạn","Nội dung"],[
    ("Pilot","SSO thử nghiệm, import danh sách lớp, rubric tự luận, dashboard phân tích"),
    ("Mở rộng","Tích hợp LMS/SIS, ngân hàng chuẩn đầu ra, kiểm thử tải và giám sát"),
    ("Sản xuất","PITR, quy trình DR, pentest, ký số biên bản, SLA và đào tạo người dùng"),
],[2100,7260])

page_break(); doc.add_heading("9. KẾT LUẬN",1)
p("UDA Assessment Hub chứng minh khả năng xây dựng một nền tảng web vừa tạo hứng thú học tập, vừa đáp ứng các nguyên tắc kiểm soát cơ bản của khảo thí đại học. Kiến trúc hiện tại đã có chuỗi kiểm duyệt đề, niêm phong, điều kiện dự thi, thi an toàn, công bố điểm, phúc khảo, đổi điểm độc lập và audit. Đây là nền tảng phù hợp để trình diễn đồ án và tiếp tục pilot có kiểm soát; chưa nên xem là hệ thống khảo thí chính thức cho tới khi hoàn tất tích hợp, đánh giá tải, an toàn thông tin và phê duyệt quy trình của nhà trường.")

doc.add_heading("TÀI LIỆU THAM KHẢO",1)
refs=[
    "Trường Đại học Đông Á. Logo Đại học Đông Á. https://donga.edu.vn/gioi-thieu/logo-dai-hoc-dong-a",
    "Trường Đại học Đông Á. Cổng thông tin và hệ sinh thái số sinh viên. https://duoc.donga.edu.vn/sinh-vien/dieu-can-biet",
    "Trường Đại học Đông Á. Hướng dẫn Canvas và đăng nhập hệ thống. https://canvas.donga.edu.vn/courses/2323/pages/canvas-va-dang-nhap-he-thong-canvas-cua-dh-dong-a",
    "Trường Đại học Đông Á. Sổ tay sinh viên và quy định đào tạo (tài liệu công khai). https://daotao.donga.edu.vn/Portals/12/document/10-08-2024/2020.2021_10_08_2024_08_01_29.pdf",
    "Next.js Documentation. https://nextjs.org/docs",
    "Supabase Documentation. Database, Auth, Row Level Security và Realtime. https://supabase.com/docs",
    "OWASP. Application Security Verification Standard. https://owasp.org/www-project-application-security-verification-standard/",
]
numbered(refs)

doc.add_heading("PHỤ LỤC A — HƯỚNG DẪN TRIỂN KHAI NHANH",1)
numbered(["Cấu hình biến môi trường Supabase/AI từ tệp mẫu.","Chạy migration theo thứ tự trong thư mục supabase/migrations.","Tạo tài khoản đầu tiên và gán vai trò admin bằng câu lệnh trong RUNBOOK.md.","Chạy npm run lint, npm test và npm run build.","Triển khai nhánh main lên Vercel; kiểm tra các luồng đăng nhập, nhập đề, QR và khảo thí."])
doc.add_heading("PHỤ LỤC B — KỊCH BẢN BẢO VỆ",1)
p("Kịch bản trình diễn 12–15 phút và bộ câu hỏi phản biện được lưu trong docs/DEMO_SCRIPT.md. Khuyến nghị chuẩn bị tối thiểu bốn tài khoản demo độc lập: giảng viên, trưởng bộ môn, khảo thí và sinh viên; thêm giám thị/đảm bảo chất lượng nếu thời gian cho phép.")

doc.core_properties.title = "Báo cáo đồ án UDA Assessment Hub"
doc.core_properties.subject = "Nền tảng học tập, nhập đề AI và quản trị khảo thí số"
doc.core_properties.author = "Nhóm phát triển UDA Assessment Hub"
doc.core_properties.keywords = "UDA, khảo thí, AI, Next.js, Supabase"
doc.save(OUT)
print(OUT.name)
